"""Mobile contact import/export: vCard (.vcf), CSV, Excel (.xlsx) and JSON.

Phone numbers are normalised to WhatsApp E.164 (`+919825011223`) and duplicates are
dropped both inside the uploaded file and against the existing HNI contact book.
"""

import csv
import io
import json
import re
from datetime import datetime

from openpyxl import Workbook, load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models.contact import (
    CATEGORIES,
    ENTERPRISE_CLASSES,
    PROJECT_TYPES,
    SECTORS,
    HniContact,
    next_inquiry_id,
)

MAX_UPLOAD_BYTES = 10 * 1024 * 1024  # 10 MB is plenty for a phone contact export

EXPORT_COLUMNS = (
    "inquiry_id",
    "client_name",
    "company_name",
    "whatsapp_number",
    "email",
    "state",
    "district",
    "taluka",
    "sector",
    "category",
    "enterprise_classification",
    "term_loan_planned",
    "project_type",
    "estimated_subsidy_lakhs",
    "tags",
    "opt_in_status",
    "assigned_ca",
)

# Longest alias first so "company name" never collapses into "name".
HEADER_ALIASES = {
    "inquiry_id": "inquiry_id",
    "inquiry": "inquiry_id",
    "inquiry no": "inquiry_id",
    "client_name": "client_name",
    "client name": "client_name",
    "customer name": "client_name",
    "contact name": "client_name",
    "full name": "client_name",
    "name": "client_name",
    "company_name": "company_name",
    "company": "company_name",
    "organisation": "company_name",
    "organization": "company_name",
    "firm": "company_name",
    "business": "company_name",
    "whatsapp_number": "whatsapp_number",
    "whatsapp number": "whatsapp_number",
    "whatsapp": "whatsapp_number",
    "mobile": "whatsapp_number",
    "phone": "whatsapp_number",
    "phone number": "whatsapp_number",
    "contact number": "whatsapp_number",
    "tel": "whatsapp_number",
    "email": "email",
    "email id": "email",
    "mail": "email",
    "state": "state",
    "district": "district",
    "city": "district",
    "taluka": "taluka",
    "tehsil": "taluka",
    "village": "taluka",
    "sector": "sector",
    "industry": "sector",
    "vertical": "sector",
    "category": "category",
    "enterprise_classification": "enterprise_classification",
    "enterprise classification": "enterprise_classification",
    "enterprise": "enterprise_classification",
    "business size": "enterprise_classification",
    "msme": "enterprise_classification",
    "term_loan_planned": "term_loan_planned",
    "term loan": "term_loan_planned",
    "term loan planned": "term_loan_planned",
    "project_type": "project_type",
    "project type": "project_type",
    "estimated_subsidy_lakhs": "estimated_subsidy_lakhs",
    "estimated subsidy": "estimated_subsidy_lakhs",
    "subsidy amount": "estimated_subsidy_lakhs",
    "subsidy": "estimated_subsidy_lakhs",
    "estimated subsidy (lakhs)": "estimated_subsidy_lakhs",
    "tags": "tags",
    "tag": "tags",
    "labels": "tags",
    "opt_in_status": "opt_in_status",
    "opt in": "opt_in_status",
    "opt-in": "opt_in_status",
    "consent": "opt_in_status",
    "assigned_ca": "assigned_ca",
    "assigned ca": "assigned_ca",
    "ca": "assigned_ca",
    "relationship manager": "assigned_ca",
}

SECTOR_ALIASES = {
    "textile": "Textile Industry",
    "textiles": "Textile Industry",
    "textile industry": "Textile Industry",
    "it": "Information Technology",
    "software": "Information Technology",
    "tech": "Information Technology",
    "information technology": "Information Technology",
    "pharma": "Pharmaceuticals",
    "pharmaceutical": "Pharmaceuticals",
    "pharmaceuticals": "Pharmaceuticals",
    "chemical": "Pharmaceuticals",
    "green": "Green Energy",
    "green energy": "Green Energy",
    "solar": "Green Energy",
    "renewable": "Green Energy",
    "energy": "Green Energy",
}

def normalize_phone(raw: object) -> str | None:
    """Any human format -> E.164, or None when it cannot be a WhatsApp number."""
    digits = re.sub(r"\D", "", str(raw or ""))
    if digits.startswith("00"):
        digits = digits[2:]
    if len(digits) == 10 and digits[0] in "6789":
        return f"+91{digits}"
    if len(digits) == 11 and digits.startswith("0"):
        return f"+91{digits[1:]}"
    if len(digits) == 12 and digits.startswith("91"):
        return f"+{digits}"
    if 11 <= len(digits) <= 15:
        return f"+{digits}"
    return None


def parse_number(value: object) -> float:
    text = str(value or "").replace(",", "").replace("₹", "").strip()
    match = re.search(r"-?\d+(?:\.\d+)?", text)
    return float(match.group()) if match else 0.0


def parse_bool(value: object, default: bool = True) -> bool:
    if value is None or value == "":
        return default
    if isinstance(value, bool):
        return value
    return str(value).strip().lower() in {"1", "true", "yes", "y", "on", "opt in", "opted in", "subscribed"}


def _normalize_header(value: object) -> str:
    return re.sub(r"[^a-z0-9]+", " ", str(value or "").strip().lower()).strip()


_ALIAS_LOOKUP = sorted(
    ((_normalize_header(alias), field) for alias, field in HEADER_ALIASES.items()),
    key=lambda row: len(row[0]),
    reverse=True,
)


def field_for_header(header: object) -> str | None:
    key = _normalize_header(header)
    if not key:
        return None
    for alias, field in _ALIAS_LOOKUP:
        if key == alias or key.startswith(alias) or key.endswith(alias):
            return field
    return None


def _canonical(value: object, options: tuple[str, ...], default: str = "") -> str:
    text = str(value or "").strip()
    if not text:
        return default
    lowered = text.lower()
    for option in options:
        if option.lower() == lowered:
            return option
    for option in options:
        if lowered in option.lower() or option.lower() in lowered:
            return option
    return text[:60]


def canonical_sector(value: object) -> str:
    text = str(value or "").strip()
    if not text:
        return ""
    alias = SECTOR_ALIASES.get(text.lower())
    return alias or _canonical(text, SECTORS)


def _record_from_mapping(mapping) -> dict:
    record: dict = {}
    for header, value in dict(mapping).items():
        field = field_for_header(header)
        if field and value not in (None, ""):
            record[field] = value
    return record


def build_contact(record: dict, *, is_mock: bool = False) -> HniContact:
    """Validate one import row / API payload into a HniContact (no session, no inquiry id)."""
    phone = normalize_phone(record.get("whatsapp_number"))
    if phone is None:
        raise ValueError("Missing or invalid WhatsApp number")

    client_name = str(record.get("client_name") or record.get("company_name") or "").strip()
    if not client_name:
        raise ValueError("Missing client name")

    return HniContact(
        inquiry_id="",
        client_name=client_name[:120],
        company_name=str(record.get("company_name") or "").strip()[:160],
        whatsapp_number=phone,
        email=str(record.get("email") or "").strip().lower()[:160],
        state=str(record.get("state") or "Gujarat").strip()[:60] or "Gujarat",
        district=str(record.get("district") or "").strip()[:80],
        taluka=str(record.get("taluka") or "").strip()[:80],
        sector=canonical_sector(record.get("sector")),
        category=_canonical(record.get("category"), CATEGORIES),
        enterprise_classification=_canonical(record.get("enterprise_classification"), ENTERPRISE_CLASSES),
        term_loan_planned="Yes" if parse_bool(record.get("term_loan_planned"), False) else "No",
        project_type=_canonical(record.get("project_type"), PROJECT_TYPES, "New"),
        estimated_subsidy_lakhs=parse_number(record.get("estimated_subsidy_lakhs")),
        tags=", ".join(dict.fromkeys(tag.strip() for tag in str(record.get("tags") or "").replace(";", ",").split(",") if tag.strip()))[:255],
        opt_in_status=parse_bool(record.get("opt_in_status"), True),
        assigned_ca=str(record.get("assigned_ca") or "").strip()[:120],
        is_mock=is_mock,
    )


# --------------------------------------------------------------------------- parsers


def parse_vcf(text: str) -> list[dict]:
    unfolded = re.sub(r"\r?\n[ \t]", "", text)
    cards: list[dict] = []
    current: dict | None = None
    for line in unfolded.splitlines():
        token = line.strip()
        upper = token.upper()
        if upper == "BEGIN:VCARD":
            current = {}
            continue
        if upper == "END:VCARD":
            if current and current.get("whatsapp_number"):
                cards.append(current)
            current = None
            continue
        if current is None or ":" not in token:
            continue
        key_part, value = token.split(":", 1)
        key = key_part.split(";")[0].strip().upper().rsplit(".", 1)[-1]
        value = value.strip()
        if not value:
            continue
        if key == "FN":
            current.setdefault("client_name", value)
        elif key == "N" and "client_name" not in current:
            parts = [part for part in value.split(";") if part.strip()]
            if parts:
                current["client_name"] = " ".join(reversed(parts[:2])).strip()
        elif key == "ORG":
            current.setdefault("company_name", value.split(";")[0])
        elif key == "TEL":
            current.setdefault("whatsapp_number", value)
        elif key == "EMAIL":
            current.setdefault("email", value)
    return cards


def parse_csv(text: str) -> list[dict]:
    reader = csv.DictReader(io.StringIO(text))
    return [record for record in (_record_from_mapping(row) for row in reader) if record]


def parse_xlsx(data: bytes) -> list[dict]:
    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    sheet = workbook.active
    rows = sheet.iter_rows(values_only=True)
    headers = next(rows, None)
    if not headers:
        return []
    records: list[dict] = []
    for row in rows:
        if row is None or all(cell in (None, "") for cell in row):
            continue
        record = _record_from_mapping(zip(headers, row))
        if record:
            records.append(record)
    workbook.close()
    return records


def parse_json(text: str) -> list[dict]:
    payload = json.loads(text)
    if isinstance(payload, dict):
        payload = payload.get("contacts") or payload.get("data") or [payload]
    if not isinstance(payload, list):
        raise ValueError("JSON must be a list of contacts or an object with a 'contacts' list")
    return [record for record in (_record_from_mapping(item) for item in payload if isinstance(item, dict)) if record]


def parse_upload(filename: str, data: bytes) -> list[dict]:
    if len(data) > MAX_UPLOAD_BYTES:
        raise ValueError("File is larger than the 10 MB import limit")
    suffix = (filename or "").rsplit(".", 1)[-1].lower()
    if suffix == "vcf":
        return parse_vcf(_decode(data))
    if suffix == "csv":
        return parse_csv(_decode(data))
    if suffix in {"xlsx", "xlsm"}:
        return parse_xlsx(data)
    if suffix == "json":
        return parse_json(_decode(data))
    raise ValueError(f"Unsupported file type '.{suffix}'. Upload a .vcf, .csv, .xlsx or .json file.")


def _decode(data: bytes) -> str:
    for encoding in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


# --------------------------------------------------------------------------- import


def import_records(db: Session, records: list[dict], *, is_mock: bool = False, auto_welcome: bool = False) -> dict:
    existing = set(db.scalars(select(HniContact.whatsapp_number)).all())
    created: list[HniContact] = []
    duplicates = 0
    invalid: list[dict] = []

    for index, record in enumerate(records, start=1):
        try:
            contact = build_contact(record, is_mock=is_mock)
        except ValueError as exc:
            invalid.append({"row": index, "reason": str(exc)})
            continue
        if contact.whatsapp_number in existing:
            duplicates += 1
            continue
        contact.inquiry_id = next_inquiry_id(db)
        db.add(contact)
        db.flush()  # so next_inquiry_id() sees the freshly imported contact
        existing.add(contact.whatsapp_number)
        created.append(contact)

    db.commit()

    welcome_sent = 0
    if auto_welcome and created:
        from app.services import whatsapp_engine

        welcome_sent = whatsapp_engine.send_welcome(db, created)

    return {
        "total_rows": len(records),
        "created": len(created),
        "duplicates": duplicates,
        "invalid": len(invalid),
        "invalid_rows": invalid[:20],
        "welcome_sent": welcome_sent,
        "contacts": [contact.public() for contact in created[:100]],
    }


# --------------------------------------------------------------------------- export


def record_for(contact: HniContact) -> dict:
    return {
        "inquiry_id": contact.inquiry_id,
        "client_name": contact.client_name,
        "company_name": contact.company_name,
        "whatsapp_number": contact.whatsapp_number,
        "email": contact.email,
        "state": contact.state,
        "district": contact.district,
        "taluka": contact.taluka,
        "sector": contact.sector,
        "category": contact.category,
        "enterprise_classification": contact.enterprise_classification,
        "term_loan_planned": contact.term_loan_planned,
        "project_type": contact.project_type,
        "estimated_subsidy_lakhs": contact.estimated_subsidy_lakhs,
        "tags": contact.tags,
        "opt_in_status": "Yes" if contact.opt_in_status else "No",
        "assigned_ca": contact.assigned_ca,
    }


def export_csv(contacts: list[HniContact]) -> str:
    buffer = io.StringIO()
    writer = csv.DictWriter(buffer, fieldnames=EXPORT_COLUMNS)
    writer.writeheader()
    for contact in contacts:
        writer.writerow(record_for(contact))
    return buffer.getvalue()


def export_xlsx(contacts: list[HniContact]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "HNI Contacts"
    sheet.append(list(EXPORT_COLUMNS))
    for contact in contacts:
        row = record_for(contact)
        sheet.append([row[column] for column in EXPORT_COLUMNS])
    buffer = io.BytesIO()
    workbook.save(buffer)
    workbook.close()
    return buffer.getvalue()


def export_vcf(contacts: list[HniContact]) -> str:
    cards = []
    for contact in contacts:
        note = f"{contact.inquiry_id} | {contact.sector or 'Sector pending'} | Subsidy Rs {contact.estimated_subsidy_lakhs:g} Lakhs"
        cards.append(
            "\r\n".join(
                [
                    "BEGIN:VCARD",
                    "VERSION:3.0",
                    f"N:;{contact.client_name};;;",
                    f"FN:{contact.client_name}",
                    f"ORG:{contact.company_name}",
                    f"TEL;TYPE=CELL:{contact.whatsapp_number}",
                    f"EMAIL:{contact.email}",
                    f"NOTE:{note}",
                    "END:VCARD",
                ]
            )
        )
    return "\r\n".join(cards) + ("\r\n" if cards else "")


def sample_csv() -> str:
    rows = [
        {
            "inquiry_id": "#123",
            "client_name": "Amit Shah",
            "company_name": "ABC Pvt Ltd",
            "whatsapp_number": "+91 98250 11223",
            "email": "amit.shah@abcgroup.in",
            "state": "Gujarat",
            "district": "Surat",
            "taluka": "Olpad",
            "sector": "Textile Industry",
            "category": "Manufacturing",
            "enterprise_classification": "MSME - Medium",
            "term_loan_planned": "Yes",
            "project_type": "Expansion",
            "estimated_subsidy_lakhs": "185.5",
            "tags": "VIP HNI, Welcome Sent",
            "opt_in_status": "Yes",
            "assigned_ca": "CA Abhishek Jain",
        },
        {
            "inquiry_id": "",
            "client_name": "Rohan Patel",
            "company_name": "Newgen InfoTech",
            "whatsapp_number": "98765 43210",
            "email": "rohan@newgeninfotech.com",
            "state": "Gujarat",
            "district": "Ahmedabad",
            "taluka": "Sanand",
            "sector": "Information Technology",
            "category": "Services",
            "enterprise_classification": "Large Enterprise",
            "term_loan_planned": "No",
            "project_type": "New",
            "estimated_subsidy_lakhs": "92",
            "tags": "High Value",
            "opt_in_status": "Yes",
            "assigned_ca": "Rahul Desai",
        },
    ]
    buffer = io.StringIO()
    writer = csv.DictWriter(buffer, fieldnames=EXPORT_COLUMNS)
    writer.writeheader()
    writer.writerows(rows)
    return buffer.getvalue()


def sample_vcf() -> str:
    return (
        "BEGIN:VCARD\r\n"
        "VERSION:3.0\r\n"
        "N:Shah;Amit;;;\r\n"
        "FN:Amit Shah\r\n"
        "ORG:ABC Pvt Ltd\r\n"
        "TEL;TYPE=CELL:+919825011223\r\n"
        "EMAIL:amit.shah@abcgroup.in\r\n"
        "NOTE:Sector: Textile Industry | Subsidy 185.5 Lakhs\r\n"
        "END:VCARD\r\n"
        "BEGIN:VCARD\r\n"
        "VERSION:3.0\r\n"
        "N:Patel;Rohan;;;\r\n"
        "FN:Rohan Patel\r\n"
        "ORG:Newgen InfoTech\r\n"
        "TEL;TYPE=CELL:+919876543210\r\n"
        "EMAIL:rohan@newgeninfotech.com\r\n"
        "NOTE:Sector: Information Technology | Subsidy 92 Lakhs\r\n"
        "END:VCARD\r\n"
    )


def export_filename(fmt: str) -> str:
    stamp = datetime.now().strftime("%Y%m%d_%H%M")
    return f"subsidyx_hni_contacts_{stamp}.{fmt}"
